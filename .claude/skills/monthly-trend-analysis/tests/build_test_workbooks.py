"""Build the portability test workbook: a deliberately DIFFERENT company and layout from the ABC Inc export.

Run:  .venv/Scripts/python tests/build_test_workbooks.py      (writes tests/test_data/bluebird-bakery.xlsx)

Differences from the ABC Inc export, each one exercising something the extractor must not assume:
  - company name on row 1, report title on row 2 (reverse order); header on row 4, not 5
  - a single label column (A), months from column B, OLDEST-first, as real date cells; 9 months only
  - a trailing "Total" column (must be ignored)
  - values only, no formulas (a formula-free export)
  - "Total for X" subtotal style (newer QBO) instead of "Total X"
  - section names "Revenue" / "Cost of Sales" / "Expenses" / "Other Income", summary row "Net Profit"
  - a group with NO header row ("Delivery Van Fuel", "Courier: Same Day", then "Total Delivery")
  - an all-blank account in the middle of the expense list ("Charitable Donations")
  - blank cells in some months; a negative contra account ("Supplier Rebates"); a negative revenue month
  - a new revenue stream that only appears in the latest month ("Wholesale Accounts")
  - accounts no category rule matches ("Mystery Shopper Program", "Instagram Ads", "Supplier Rebates")
  - ONE deliberately wrong subtotal: "Total for Professional Services" overstated by 25.00 in Mar 2026
  - Balance Sheet: labels in column B, header "Account" in B, month headers as "Sep-26" text, newest-first
"""
from datetime import datetime
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "test_data" / "bluebird-bakery.xlsx"
MONTHS = [datetime(2026, m, 1) for m in range(1, 10)]          # Jan..Sep 2026, oldest first
N = len(MONTHS)


def series(base, step=0.0, blanks=()):
    return [None if i in blanks else round(base + step * i, 2) for i in range(N)]


REVENUE = {
    "Bread Sales": series(42000, 900),
    "Pastry Sales": series(18500, -350),
    "Catering Revenue": [6000, 5500, 0, 7200, 6100, 8800, 9100, 4000, 3800],
    "Gift Cards": [500, 450, -300, 600, 550, 500, 650, 700, 600],          # negative month = refunds
    "Wholesale Accounts": [None] * 8 + [12000],                              # new stream in the latest month
}
COGS = {
    "Flour & Ingredients": series(14000, 250),
    "Packaging": series(2100, 40),
    "Kitchen Wages": series(9800, 0),
}
EXPENSES_TOP = {
    "Wages & Salaries": series(16000, 0) [:7] + [16000, 21000],              # jump in the latest month
    "Payroll Taxes": series(1600, 0),
}
SHOP = {"Rent - High Street": series(4500, 0), "Utilities": series(820, 15), "Equipment Repairs": series(300, 0, blanks=(1, 4))}
DELIVERY = {"Delivery Van Fuel": series(640, 10), "Courier: Same Day": [210, 180, 260, 0, 190, 230, 250, 300, 310]}
MIDDLE = {"Instagram Ads": series(950, 50, blanks=(2,)), "Mystery Shopper Program": [0, 0, 400, 0, 0, 400, 0, 0, 400]}
PROF = {"Bookkeeping": series(700, 0), "Legal": [0, 0, 1500, 0, 0, 0, 0, 0, 0]}
TAIL = {"Supplier Rebates": [-250, -250, -300, -250, -250, -400, -250, -250, -250]}
OTHER_INCOME = {"Interest Earned": series(35, 1)}


def v(x):
    return 0.0 if x is None else x


def colsum(*groups):
    return [round(sum(v(vals[i]) for g in groups for vals in g.values()), 2) for i in range(N)]


def write_pl(ws):
    ws["A1"], ws["A2"], ws["A3"] = "Bluebird Bakery Co", "Profit and Loss", "January - September, 2026"
    ws.cell(4, 1, "Account")
    for j, m in enumerate(MONTHS):
        ws.cell(4, 2 + j, m).number_format = "mmm yyyy"
    ws.cell(4, 2 + N, "Total")
    r = 5

    def put(label, values=None, total_col=True):
        nonlocal r
        ws.cell(r, 1, label)
        if values is not None:
            for j, x in enumerate(values):
                if x is not None:
                    ws.cell(r, 2 + j, x)
            if total_col:
                ws.cell(r, 2 + N, round(sum(v(x) for x in values), 2))
        r += 1

    put("Revenue")
    for k, vals in REVENUE.items():
        put(k, vals)
    rev = colsum(REVENUE)
    put("Total for Revenue", rev)
    put("Cost of Sales")
    for k, vals in COGS.items():
        put(k, vals)
    cogs = colsum(COGS)
    put("Total for Cost of Sales", cogs)
    put("Gross Profit", [round(a - b, 2) for a, b in zip(rev, cogs)])
    put("Expenses")
    for k, vals in EXPENSES_TOP.items():
        put(k, vals)
    put("Shop Costs")
    for k, vals in SHOP.items():
        put(k, vals)
    put("Total for Shop Costs", colsum(SHOP))
    for k, vals in DELIVERY.items():                                  # no header row for this group
        put(k, vals)
    put("Total Delivery", colsum(DELIVERY))
    put("Charitable Donations")                                       # blank in every month, mid-list
    for k, vals in MIDDLE.items():
        put(k, vals)
    put("Professional Services")
    for k, vals in PROF.items():
        put(k, vals)
    prof_stated = colsum(PROF)
    prof_stated[2] = round(prof_stated[2] + 25.0, 2)                  # deliberate data-quality error (Mar 2026)
    put("Total for Professional Services", prof_stated)
    for k, vals in TAIL.items():
        put(k, vals)
    expenses = colsum(EXPENSES_TOP, SHOP, DELIVERY, MIDDLE, PROF, TAIL)  # true detail sum
    put("Total for Expenses", expenses)
    put("Other Income")
    for k, vals in OTHER_INCOME.items():
        put(k, vals)
    oi = colsum(OTHER_INCOME)
    put("Total for Other Income", oi)
    put("Net Profit", [round(a - b - c + d, 2) for a, b, c, d in zip(rev, cogs, expenses, oi)])


def write_bs(ws):
    ws["B1"], ws["B2"], ws["B3"] = "Bluebird Bakery Co", "Balance Sheet", "As of September 30, 2026"
    ws.cell(5, 2, "Account")
    order = list(reversed(range(N)))                                   # newest first
    for j, i in enumerate(order):
        ws.cell(5, 3 + j, MONTHS[i].strftime("%b-%y"))
    checking = [52000 - 1800 * i for i in range(N)]
    receivable = [9000 + 300 * i for i in range(N)]
    card = [3100 + 120 * i for i in range(N)]
    retained = [40000] * N
    current = [checking[i] + receivable[i] - card[i] - retained[i] for i in range(N)]
    rows = [
        ("Assets", None), ("Bank Accounts", None), ("Business Checking", checking), ("Total for Bank Accounts", checking),
        ("Accounts Receivable", receivable),
        ("Total for Assets", [a + b for a, b in zip(checking, receivable)]),
        ("Liabilities and Equity", None), ("Liabilities", None), ("Credit Card", card), ("Total for Liabilities", card),
        ("Equity", None), ("Retained Earnings", retained), ("Current Year Earnings", current),
        ("Total for Equity", [a + b for a, b in zip(retained, current)]),
        ("Total for Liabilities and Equity", [card[i] + retained[i] + current[i] for i in range(N)]),
    ]
    for k, (label, vals) in enumerate(rows):
        ws.cell(6 + k, 2, label)
        if vals:
            for j, i in enumerate(order):
                ws.cell(6 + k, 3 + j, vals[i])


def main():
    wb = openpyxl.Workbook()
    write_pl(wb.active)
    wb.active.title = "Profit and Loss"
    write_bs(wb.create_sheet("Balance Sheet"))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
