"""Tests for the monthly-trend-analysis skill.

Run from the skill folder:  .venv/Scripts/python -m unittest discover -s tests -v   (POSIX: .venv/bin/python)
The real-workbook tests are skipped if data/*.xlsx is not present.
"""
import json
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))

from trends import analysis, mapping, workbook  # noqa: E402
import run_trends  # noqa: E402

CONFIG = mapping.load_config()
REPO_DATA = SKILL.parents[2] / "data"
ABC = next(iter(sorted(REPO_DATA.glob("ABC Inc*.xlsx"))), None) if REPO_DATA.is_dir() else None
BLUEBIRD = SKILL / "tests" / "test_data" / "bluebird-bakery.xlsx"


def sheet_from_rows(rows, months=("Jan 2026", "Feb 2026", "Mar 2026"), formulas=None):
    """Build (values_ws, formulas_ws) in memory. rows: [(label, [values] | None)]. formulas: {row_index: str}."""
    wv, wf = openpyxl.Workbook().active, openpyxl.Workbook().active
    for ws in (wv, wf):
        ws.cell(1, 1, "Test Co")
        ws.cell(3, 1, "Account")
        for j, m in enumerate(months):
            ws.cell(3, 2 + j, m)
    for i, (label, vals) in enumerate(rows):
        r = 4 + i
        for ws in (wv, wf):
            ws.cell(r, 1, label)
        if vals:
            for j, x in enumerate(vals):
                wv.cell(r, 2 + j, x)
                wf.cell(r, 2 + j, (formulas or {}).get(i, x) if formulas and i in formulas else x)
    return wv, wf


def account_names(sheet):
    return [a.name for s in sheet.sections for a, _ in s.accounts()]


class UnitTests(unittest.TestCase):
    def test_month_formats(self):
        for text in ("Jun 2026", "30 Jun 2026", "June 2026", "Jun-26", "2026-06", datetime(2026, 6, 1)):
            self.assertEqual(workbook.parse_month(text), (date(2026, 6, 1), False), text)
        self.assertEqual(workbook.parse_month("Oct 1-12, 2026"), (date(2026, 10, 1), True))
        for not_month in ("Account", "Total", None, 123):
            self.assertEqual(workbook.parse_month(not_month)[0], None)

    def test_total_label_matching(self):
        self.assertEqual(workbook.total_target("Total S & M Payroll"), "S & M Payroll")
        self.assertEqual(workbook.total_target("Total for Revenue"), "Revenue")
        self.assertIsNone(workbook.total_target("Totally Awesome Revenue"))
        self.assertEqual(workbook.norm("Cost of Goods Sold:"), workbook.norm("cost of goods  sold"))

    def test_numbers(self):
        self.assertEqual(workbook.parse_number("(1,234.50)"), -1234.5)
        self.assertEqual(workbook.parse_number("$2,000"), 2000.0)
        self.assertIsNone(workbook.parse_number(""))
        self.assertIsNone(workbook.parse_number(None))

    def test_categorisation_rules(self):
        cat = lambda name, groups=(): mapping.opex_category(name, list(groups), CONFIG)[0]
        self.assertEqual(cat("Rent"), "Facilities & Office")                 # regression: \\b escaping in JSON
        self.assertEqual(cat("Marketing - Payroll", ["S & M Payroll"]), "Payroll & Benefits")
        self.assertEqual(cat("Events & Conferences", ["Marketing"]), "Sales & Marketing")
        self.assertEqual(cat("Marketing/Advertising: Website"), "Sales & Marketing")
        self.assertEqual(cat("Income Tax Benefit"), "Tax & Interest (non-operating)")
        self.assertTrue(mapping.opex_category("Interest Income", [], CONFIG)[1])
        self.assertEqual(cat("Completely Unknown Thing"), CONFIG["opex_categories"]["fallback_category"])

    def test_section_roles(self):
        role = lambda n: mapping.section_role(n, CONFIG)
        self.assertEqual(role("Income"), "revenue")
        self.assertEqual(role("Revenue"), "revenue")
        self.assertEqual(role("Cost of Goods Sold"), "cogs")
        self.assertEqual(role("Cost of Sales"), "cogs")
        self.assertEqual(role("Operating Expenses"), "operating_expenses")
        self.assertEqual(role("Expenses"), "operating_expenses")
        self.assertEqual(role("Other Income"), "other_income")
        self.assertEqual(role("Other Expenses"), "other_expense")


class SyntheticSheetTests(unittest.TestCase):
    def test_nested_subtotal_is_not_double_counted(self):
        wv, wf = sheet_from_rows([
            ("Expenses", None),
            ("Rent", [100, 100, 100]),
            ("Payroll", None),
            ("Wages", [50, 60, 70]),
            ("Taxes", [5, 6, 7]),
            ("Total Payroll", [55, 66, 77]),
            ("Total Expenses", [155, 166, 177]),
        ])
        s = workbook.parse_sheet(wv, wf)
        self.assertEqual(account_names(s), ["Rent", "Wages", "Taxes"])
        self.assertEqual(s.sections[0].detail_sum(3), [155, 166, 177])
        self.assertEqual(s.warnings, [])

    def test_mismatch_is_flagged_not_trusted(self):
        wv, wf = sheet_from_rows([("Revenue", None), ("A", [10, 10, 10]), ("B", [5, 5, 5]),
                                  ("Total Revenue", [15, 99, 15])])
        s = workbook.parse_sheet(wv, wf)
        mism = [w for w in s.warnings if w["type"] == "subtotal_mismatch"]
        self.assertEqual(len(mism), 1)
        self.assertEqual((mism[0]["month"], mism[0]["stated"], mism[0]["detail_sum"]), ("2026-02", 99, 15))
        self.assertEqual(s.sections[0].detail_sum(3), [15, 15, 15])   # detail rows drive the analysis

    def test_total_without_header_is_inferred_from_numbers(self):
        wv, wf = sheet_from_rows([("Expenses", None), ("Rent", [100, 100, 100]), ("Fuel", [10, 20, 30]),
                                  ("Courier", [1, 2, 3]), ("Total Delivery", [11, 22, 33]),
                                  ("Total Expenses", [111, 122, 133])])
        s = workbook.parse_sheet(wv, wf)
        groups = [c for c in s.sections[0].children if isinstance(c, workbook.Section)]
        self.assertEqual([(g.name, g.inferred, [a.name for a in g.children]) for g in groups],
                         [("Delivery", True, ["Fuel", "Courier"])])
        self.assertTrue(any(w["type"] == "inferred_group" for w in s.warnings))
        self.assertFalse(any(w["type"] == "subtotal_mismatch" for w in s.warnings))

    def test_unreconcilable_total_is_excluded_and_flagged(self):
        wv, wf = sheet_from_rows([("Expenses", None), ("Rent", [100, 100, 100]), ("Total Mystery", [7, 7, 7]),
                                  ("Total Expenses", [100, 100, 100])])
        s = workbook.parse_sheet(wv, wf)
        self.assertEqual(account_names(s), ["Rent"])
        self.assertTrue(any(w["type"] == "unmatched_total" for w in s.warnings))

    def test_blank_account_mid_list_does_not_capture_following_rows(self):
        wv, wf = sheet_from_rows([("Expenses", None), ("Donations", None), ("Ads", [1, 2, 3]),
                                  ("Total Expenses", [1, 2, 3])])
        s = workbook.parse_sheet(wv, wf)
        top = s.sections[0]
        self.assertEqual([c.name for c in top.children], ["Donations", "Ads"])   # siblings, not nested
        self.assertTrue(all(isinstance(c, workbook.Account) for c in top.children))
        self.assertTrue(any(w["type"] == "missing_total" for w in s.warnings))

    def test_blank_cells_are_zero_and_counted(self):
        wv, wf = sheet_from_rows([("Revenue", None), ("A", [10, None, 30]), ("Total Revenue", [10, 0, 30])])
        s = workbook.parse_sheet(wv, wf)
        acc = next(s.sections[0].accounts())[0]
        self.assertEqual((acc.values, acc.blank_months), ([10, 0, 30], 1))

    def test_formula_in_detail_row_is_flagged(self):
        wv, wf = sheet_from_rows([("Revenue", None), ("A", [10, 10, 10]), ("Sneaky Sum", [10, 10, 10]),
                                  ("Total Revenue", [20, 20, 20])], formulas={2: "=B5"})
        s = workbook.parse_sheet(wv, wf)
        self.assertTrue(any(w["type"] == "formula_in_detail" and w["label"] == "Sneaky Sum" for w in s.warnings))

    def test_newest_first_columns_are_sorted(self):
        wv, wf = sheet_from_rows([("Revenue", None), ("A", [3, 2, 1]), ("Total Revenue", [3, 2, 1])],
                                 months=("Mar 2026", "Feb 2026", "Jan 2026"))
        s = workbook.parse_sheet(wv, wf)
        self.assertEqual(s.months, [date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)])
        self.assertEqual(next(s.sections[0].accounts())[0].values, [1, 2, 3])


def run_on(path):
    v, f = workbook.load_workbook_pair(path)
    is_ws = workbook.find_sheet(v, CONFIG["sheets"]["income_statement"])
    bs_ws = workbook.find_sheet(v, CONFIG["sheets"]["balance_sheet"])
    income = workbook.parse_sheet(is_ws, f[is_ws.title])
    balance = workbook.parse_sheet(bs_ws, f[bs_ws.title]) if bs_ws else None
    return income, balance, analysis.analyze(income, balance, CONFIG, path)


@unittest.skipUnless(ABC and ABC.exists(), "ABC Inc workbook not present in data/")
class RealWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.income, cls.balance, cls.a = run_on(ABC)
        cls.raw = openpyxl.load_workbook(ABC, data_only=True)

    def raw_row(self, sheet, label, first_col):
        ws = self.raw[sheet]
        r = next(r for r in range(1, ws.max_row + 1) for c in (1, 2, 3) if ws.cell(r, c).value == label)
        return list(reversed([ws.cell(r, c).value or 0 for c in range(first_col, first_col + 18)]))

    def test_structure(self):
        self.assertEqual(len(self.income.months), 18)
        self.assertEqual([s.name for s in self.income.sections], ["Income", "Cost of Goods Sold", "Operating Expenses"])
        self.assertEqual(len(self.a["revenue"]["streams"]), 6)
        self.assertEqual(self.income.warnings, [])
        self.assertEqual(self.balance.warnings, [])

    def test_no_subtotal_or_header_rows_counted_as_accounts(self):
        names = account_names(self.income)
        self.assertFalse([n for n in names if n.lower().startswith("total")])
        for header in ("S & M Payroll", "Marketing", "Income", "Operating Expenses", "Gross Profit"):
            self.assertNotIn(header, names)
        self.assertEqual(len(names), 6 + 7 + 49 + 11 + 10)

    def test_figures_match_raw_sheet_rows(self):
        pairs = [("Total Income", self.a["revenue"]["total"]["monthly"]),
                 ("Gross Profit", [m["gross_profit"] for m in self.a["gross_margin"]["monthly"]]),
                 ("Total Operating Expenses", self.a["operating_expenses"]["total"]["monthly"]),
                 ("Net Income", self.a["net_income"]["monthly"])]
        for label, mine in pairs:
            raw = self.raw_row("Income Statement", label, 3)
            self.assertLess(max(abs(x - y) for x, y in zip(raw, mine)), 0.01, label)
        cash = self.raw_row("Balance Sheet", "Total Cash and Cash Equivalents", 4)
        self.assertLess(max(abs(x - y) for x, y in zip(cash, self.a["balance_sheet"]["lines"]["cash"]["monthly"])), 0.01)

    def test_ttm_mom_and_yoy_math_matches_raw_cells(self):
        # Added in Task 3: the suite checked monthly series but never asserted TTM / YoY values themselves.
        rev = self.raw_row("Income Statement", "Total Income", 3)          # oldest-first after reversal
        gp = self.raw_row("Income Statement", "Gross Profit", 3)
        ni = self.raw_row("Income Statement", "Net Income", 3)
        t = self.a["revenue"]["total"]
        self.assertAlmostEqual(t["ttm"], sum(rev[-12:]), places=2)
        self.assertAlmostEqual(t["ttm_prior_month"], sum(rev[-13:-1]), places=2)
        self.assertAlmostEqual(t["mom_pct"], (rev[-1] - rev[-2]) / abs(rev[-2]) * 100, places=3)
        self.assertAlmostEqual(t["yoy_pct"], (rev[-1] - rev[-13]) / abs(rev[-13]) * 100, places=3)
        self.assertAlmostEqual(self.a["gross_margin"]["ttm_gm_pct"], sum(gp[-12:]) / sum(rev[-12:]) * 100, places=3)
        self.assertAlmostEqual(self.a["net_income"]["ttm"], sum(ni[-12:]), places=2)
        self.assertIsNone(self.a["revenue"]["total"]["ttm_monthly"][10])     # 11th month: not yet 12 months
        self.assertIsNotNone(self.a["revenue"]["total"]["ttm_monthly"][11])

    def test_vanished_line_wording_says_zero_since(self):
        # Test 29 (intentional behaviour change, DECISIONS_LOG D-45): a line that was already zero before the latest
        # month must say "has been zero since <month>" with its last non-zero month, not "is zero in Jun 2026".
        flags = [f["message"] for f in self.a["flags"] if f["kind"] == "vanished_expense_line"]
        text = "\n".join(flags)
        self.assertIn("Expense line 'Advisory Services' (Professional Fees) has been zero since Feb 2026 "
                      "(last non-zero: Jan 2026, 2,500); it averaged 1,250", text)
        self.assertIn("Expense line 'Professional Fees' (Professional Fees) has been zero since Apr 2026 "
                      "(last non-zero: Mar 2026, 3,773); it averaged 1,045", text)
        # Lines that really went to zero in the latest month keep the original wording.
        self.assertIn("Expense line 'Federal & State Taxes' (Finance & Admin) is zero in Jun 2026; it averaged 999", text)
        self.assertIn("Expense line 'Marketing/Advertising: Miscellaneous' (Sales & Marketing) is zero in Jun 2026; "
                      "it averaged 721", text)
        self.assertNotIn("'Advisory Services' (Professional Fees) is zero in Jun 2026", text)
        self.assertEqual(len(flags), 4)

    def test_vanished_revenue_stream_wording_says_zero_since(self):
        # Test 30 (intentional behaviour change, DECISIONS_LOG D-48): same "zero since" logic for revenue streams.
        text = "\n".join(f["message"] for f in self.a["flags"] if f["kind"] == "revenue_stream")
        self.assertIn("Revenue stream 'SKU Profiling' has been zero since May 2026 (last non-zero: Apr 2026, 10,000) "
                      "after 10,000 at its peak in the prior 6 months.", text)
        self.assertNotIn("Revenue stream 'SKU Profiling' is zero in Jun 2026", text)
        # Streams that really stopped in Jun 2026 keep the original wording, byte for byte.
        self.assertIn("Revenue stream 'Partnership Income' is zero in Jun 2026 after 1,608 at its peak in the prior 6 months.", text)
        self.assertIn("Revenue stream 'Professional Services' is zero in Jun 2026 after 15,833 at its peak in the prior 6 months.",
                      text)

    def test_all_cross_checks_pass(self):
        self.assertEqual({c["status"] for c in self.a["data_quality"]["checks"]}, {"pass"})

    def test_key_flags(self):
        text = " ".join(f["message"] for f in self.a["flags"])
        self.assertIn("-6.0% vs May 2026", text)                          # Jun 2026 revenue drop
        self.assertIn("negative for 6 consecutive months", text)
        self.assertIn("Dec 2025 shows positive net income", text)          # driven by Income Tax Benefit
        self.assertIn("Income Tax Benefit", text)
        self.assertEqual(self.a["operating_expenses"]["fallback_accounts"], ["General Expenses"])


@unittest.skipUnless(BLUEBIRD.exists(), "run tests/build_test_workbooks.py first")
class PortabilityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.income, cls.balance, cls.a = run_on(BLUEBIRD)

    def test_different_company_and_layout(self):
        m = self.a["meta"]
        self.assertEqual(m["company"], "Bluebird Bakery Co")
        self.assertEqual(m["months"][0], "2026-01")
        self.assertEqual(len(m["months"]), 9)                    # trailing "Total" column ignored
        self.assertFalse(m["ttm_available"])
        self.assertEqual([s["name"] for s in self.a["revenue"]["streams"]],
                         ["Bread Sales", "Pastry Sales", "Catering Revenue", "Gift Cards", "Wholesale Accounts"])

    def test_structural_variants_handled(self):
        types = {w["type"] for w in self.income.warnings}
        self.assertTrue({"inferred_group", "missing_total", "subtotal_mismatch"} <= types)
        mism = [w for w in self.income.warnings if w["type"] == "subtotal_mismatch"]
        self.assertEqual([(w["month"], w["difference"]) for w in mism], [("2026-03", 25.0)])

    def test_totals_reconcile_with_workbook_summary_rows(self):
        status = {c["check"].split(" ")[0]: c["status"] for c in self.a["data_quality"]["checks"]}
        self.assertEqual(status["gross_profit"], "pass")
        self.assertEqual(status["net_income"], "pass")          # includes the Other Income section
        self.assertEqual(status["operating_income"], "skipped")  # no such row in this export
        self.assertTrue(any("balance sheet:" in c["check"] and c["status"] == "pass" for c in self.a["data_quality"]["checks"]))

    def test_unmapped_accounts_fall_into_other(self):
        self.assertEqual(sorted(self.a["operating_expenses"]["fallback_accounts"]),
                         ["Instagram Ads", "Mystery Shopper Program", "Supplier Rebates"])

    def test_percent_rounding_is_consistent_between_table_and_flags(self):
        # Regression: 59.551 was stored as 59.55 and displayed 59.5% in the table but 59.6% in the flag.
        aug = self.a["gross_margin"]["monthly"][-2]["gm_pct"]
        text = " ".join(f["message"] for f in self.a["flags"] if f["kind"] == "gross_margin")
        self.assertIn(f"{aug:.1f}% in Aug 2026", text)

    def test_prior_month_zero_is_not_labelled_new(self):
        # Regression: Mystery Shopper Program (quarterly, 400 in Mar/Jun/Sep) was shown as "new" in Sep.
        from trends import report
        acct = next(x for x in self.a["operating_expenses"]["accounts"] if x["name"] == "Mystery Shopper Program")
        self.assertEqual(report.mom_cell(acct), "from 0")
        self.assertNotIn("New expense line 'Mystery Shopper Program'", " ".join(f["message"] for f in self.a["flags"]))

    def test_flags_are_sensible(self):
        text = " ".join(f["message"] for f in self.a["flags"])
        self.assertIn("New revenue stream 'Wholesale Accounts'", text)
        self.assertNotIn("Single-month spike: 'Wholesale Accounts'", text)   # not double-reported
        self.assertIn("'Gift Cards' is negative in Mar 2026", text)
        self.assertIn("'Bank Accounts'", text)                                # 6-month cash decline


class CliTests(unittest.TestCase):
    def test_cli_writes_report_and_json(self):
        with tempfile.TemporaryDirectory() as out:
            self.assertEqual(run_trends.main([str(BLUEBIRD), "--out-dir", out, "--months", "3"]), 0)
            files = sorted(p.name for p in Path(out).iterdir())
            self.assertEqual(files, ["bluebird-bakery-co-2026-09-trends.json", "bluebird-bakery-co-2026-09-trends.md"])
            md = (Path(out) / files[1]).read_text(encoding="utf-8")
            self.assertIn("<!-- NARRATIVE:", md)
            self.assertIn("DRAFT for reviewer", md)
            json.loads((Path(out) / files[0]).read_text(encoding="utf-8"))

    def test_missing_workbook_is_a_clean_error(self):
        self.assertEqual(run_trends.main(["does-not-exist.xlsx", "--out-dir", tempfile.gettempdir()]), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
